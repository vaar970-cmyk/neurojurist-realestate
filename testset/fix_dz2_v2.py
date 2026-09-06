# -*- coding: utf-8 -*-
import docx

path = r"C:\Users\Варя\Documents\Myprojects\Дипломная работа\testset\ДЗ2.docx"
d = docx.Document(path)
paras = d.paragraphs

def set_text(p, new_text):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.add_run(new_text)

replacements = {
    'ГК РФ —': 'ГК РФ (garant.ru) —',
    'ЖК РФ —': 'ЖК РФ (consultant.ru) —',
    'СК РФ —': 'СК РФ (consultant.ru) —',
    'ФЗ № 218-ФЗ': 'ФЗ № 218-ФЗ (garant.ru)',
    'ФЗ № 102-ФЗ': 'ФЗ № 102-ФЗ (garant.ru)',
    'ФЗ № 214-ФЗ': 'ФЗ № 214-ФЗ (garant.ru)',
    'НК РФ —': 'НК РФ (garant.ru) —',
}

for p in paras:
    t = p.text
    # ФЗ-218: rewrite with concrete nouns instead of dangling words
    if t.startswith('ФЗ № 218-ФЗ'):
        set_text(p, 'ФЗ № 218-ФЗ (garant.ru) «О государственной регистрации недвижимости» — порядок государственной регистрации прав на недвижимость: основания и сроки регистрации, основания и сроки приостановления регистрации, порядок исправления ошибок в Едином государственном реестре недвижимости (ЕГРН), ответственность регистратора за нарушения при регистрации. Не включены главы о структуре и полномочиях органов регистрации и об информационном обмене между госорганами — это не относится к сделкам физлиц.')
        continue
    if t.startswith('ГК РФ —'):
        set_text(p, t.replace('ГК РФ —', 'ГК РФ (garant.ru) —', 1))
        continue
    if t.startswith('ЖК РФ —'):
        set_text(p, t.replace('ЖК РФ —', 'ЖК РФ (consultant.ru) —', 1))
        continue
    if t.startswith('СК РФ —'):
        set_text(p, t.replace('СК РФ —', 'СК РФ (consultant.ru) —', 1))
        continue
    if t.startswith('ФЗ № 102-ФЗ'):
        set_text(p, t.replace('ФЗ № 102-ФЗ', 'ФЗ № 102-ФЗ (garant.ru)', 1))
        continue
    if t.startswith('ФЗ № 214-ФЗ'):
        set_text(p, t.replace('ФЗ № 214-ФЗ', 'ФЗ № 214-ФЗ (garant.ru)', 1))
        continue
    if t.startswith('НК РФ —'):
        set_text(p, t.replace('НК РФ —', 'НК РФ (garant.ru) —', 1))
        continue
    if t.startswith('Отобраны 10 постановлений'):
        set_text(p, 'Отобраны 10 постановлений Пленума Верховного Суда РФ по темам сделок с жилой недвижимостью — целевым поиском по каждой теме периметра (не сплошным просмотром всех постановлений Пленума), с проверкой действующего статуса по истории редакций документа. Источники: 1 документ — garant.ru, 9 документов — consultant.ru.')
        continue
    if t.startswith('Судебная практика по конкретным статьям'):
        set_text(p, t + ' Источник — gkrfkod.ru.')
        continue
    if t.startswith('Отобраны ежеквартальные обзоры'):
        set_text(p, t.strip() + ' Источник — consultant.ru.')
        continue

d.save(path)
print("done")
