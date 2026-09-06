import docx
import sys

def extract(path, out_path):
    d = docx.Document(path)
    with open(out_path, "w", encoding="utf-8") as f:
        for para in d.paragraphs:
            f.write(para.text + "\n")
        for table in d.tables:
            for row in table.rows:
                f.write(" | ".join(c.text for c in row.cells) + "\n")

if __name__ == "__main__":
    path, out_path = sys.argv[1], sys.argv[2]
    extract(path, out_path)
    print("done")
